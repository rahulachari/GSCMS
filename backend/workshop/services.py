import os
import io
import zipfile
import shutil
from datetime import datetime, date
from decimal import Decimal
from django.conf import settings
from django.core import serializers
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import (
    Customer, Order, Job, JobStageHistory, JobPhoto, Worker,
    GoldTransaction, Stone, StoneTransaction, Payment, Expense,
    Delivery, QualityCheck, AuditLog, BackupRecord, Notification
)

def log_audit(user, action, instance, details="", request=None):
    """Immutable audit logging for all critical workshop operations"""
    try:
        ip = None
        if request:
            x_forwarded = request.META.get('HTTP_X_FORWARDED_FOR')
            if x_forwarded:
                ip = x_forwarded.split(',')[0].strip()
            else:
                ip = request.META.get('REMOTE_ADDR')

        model_name = instance.__class__.__name__ if instance else "System"
        obj_id = str(getattr(instance, 'pk', '')) if instance else ""
        obj_repr = str(instance)[:250] if instance else "N/A"

        AuditLog.objects.create(
            user=user if getattr(user, 'is_authenticated', False) else None,
            action=action,
            model_name=model_name,
            object_id=obj_id,
            object_repr=obj_repr,
            details=details,
            ip_address=ip
        )
    except Exception as e:
        print(f"Error logging audit: {e}")


def apply_apple_excel_styling(ws, title, columns):
    """Format openpyxl worksheet with elegant, clean Apple-grade typography and table headers"""
    # Title Row
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    title_cell = ws.cell(row=1, column=1, value=title.upper())
    title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="1D1D1F")
    title_cell.fill = PatternFill(start_color="F5F5F7", end_color="F5F5F7", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 36

    # Generated Date Row
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columns))
    meta_cell = ws.cell(row=2, column=1, value=f"Generated on {timezone.now().strftime('%d %B %Y, %I:%M %p')} | Sri Swarna Jewellers Workshop Management System")
    meta_cell.font = Font(name="Segoe UI", size=9, color="707070", italic=True)
    meta_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 20

    # Header Row
    header_fill = PatternFill(start_color="1D1D1F", end_color="1D1D1F", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='E5E5EA'),
        right=Side(style='thin', color='E5E5EA'),
        top=Side(style='thin', color='E5E5EA'),
        bottom=Side(style='thin', color='E5E5EA')
    )

    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=3, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[3].height = 28
    ws.freeze_panes = "A4"


def export_jobs_to_excel(jobs_queryset=None):
    if jobs_queryset is None:
        jobs_queryset = Job.objects.all().select_related('order', 'order__customer', 'current_worker')

    wb = Workbook()
    ws = wb.active
    ws.title = "Workshop Jobs"

    columns = [
        "Job ID", "Order No.", "Customer", "Jewellery Item", "Qty",
        "Purity", "Est. Wt (g)", "Final Gross (g)", "Net Gold (g)", "Stone Wt (g)",
        "Current Stage", "Worker", "Status", "Priority", "Target Date", "Is Delayed?"
    ]
    apply_apple_excel_styling(ws, "Master Jewellery Production Jobs", columns)

    row_idx = 4
    data_font = Font(name="Segoe UI", size=10, color="1D1D1F")
    border_side = Side(style='thin', color='E5E5EA')
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    for job in jobs_queryset:
        ws.cell(row=row_idx, column=1, value=job.job_id)
        ws.cell(row=row_idx, column=2, value=job.order.order_number)
        ws.cell(row=row_idx, column=3, value=job.order.customer.name)
        ws.cell(row=row_idx, column=4, value=job.title)
        ws.cell(row=row_idx, column=5, value=job.quantity)
        ws.cell(row=row_idx, column=6, value=job.gold_purity)
        
        c7 = ws.cell(row=row_idx, column=7, value=float(job.expected_weight))
        c7.number_format = '0.000'
        
        c8 = ws.cell(row=row_idx, column=8, value=float(job.final_gross_weight or 0))
        c8.number_format = '0.000'
        
        c9 = ws.cell(row=row_idx, column=9, value=float(job.net_gold_weight or 0))
        c9.number_format = '0.000'

        c10 = ws.cell(row=row_idx, column=10, value=float(job.stone_weight or 0))
        c10.number_format = '0.000'

        ws.cell(row=row_idx, column=11, value=job.get_current_stage_display())
        ws.cell(row=row_idx, column=12, value=job.current_worker.name if job.current_worker else "Unassigned")
        ws.cell(row=row_idx, column=13, value=job.get_overall_status_display())
        ws.cell(row=row_idx, column=14, value=job.get_priority_display())
        ws.cell(row=row_idx, column=15, value=job.target_date.strftime("%d/%m/%Y %H:%M") if job.target_date else "-")
        ws.cell(row=row_idx, column=16, value="YES" if job.is_delayed else "NO")

        # Zebra striping
        row_fill = PatternFill(start_color="FFFFFF" if row_idx % 2 == 0 else "FBFBFD", fill_type="solid")
        for col_idx in range(1, len(columns) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.fill = row_fill
            cell.border = cell_border
            if col_idx in [5, 6, 11, 12, 13, 14, 15, 16]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        
        ws.row_dimensions[row_idx].height = 22
        row_idx += 1

    # Auto-adjust column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def export_gold_ledger_to_excel(transactions=None):
    if transactions is None:
        transactions = GoldTransaction.objects.all().select_related('job', 'worker', 'created_by')

    wb = Workbook()
    ws = wb.active
    ws.title = "Gold Ledger"

    columns = [
        "Txn ID", "Date / Time", "Job ID", "Worker", "Transaction Type",
        "Purity", "Weight (g)", "Balance After (g)", "Recorded By", "Notes"
    ]
    apply_apple_excel_styling(ws, "Master Gold Accountability Ledger", columns)

    row_idx = 4
    data_font = Font(name="Segoe UI", size=10, color="1D1D1F")
    border_side = Side(style='thin', color='E5E5EA')
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    for txn in transactions:
        ws.cell(row=row_idx, column=1, value=txn.transaction_id)
        ws.cell(row=row_idx, column=2, value=txn.created_at.strftime("%d/%m/%Y %H:%M"))
        ws.cell(row=row_idx, column=3, value=txn.job.job_id if txn.job else "-")
        ws.cell(row=row_idx, column=4, value=txn.worker.name if txn.worker else "-")
        ws.cell(row=row_idx, column=5, value=txn.get_transaction_type_display())
        ws.cell(row=row_idx, column=6, value=txn.purity)

        c7 = ws.cell(row=row_idx, column=7, value=float(txn.weight_grams))
        c7.number_format = '0.000'

        c8 = ws.cell(row=row_idx, column=8, value=float(txn.balance_after_grams))
        c8.number_format = '0.000'

        ws.cell(row=row_idx, column=9, value=txn.created_by.get_full_name() or txn.created_by.username if txn.created_by else "-")
        ws.cell(row=row_idx, column=10, value=txn.notes or "-")

        row_fill = PatternFill(start_color="FFFFFF" if row_idx % 2 == 0 else "FBFBFD", fill_type="solid")
        for col_idx in range(1, len(columns) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.fill = row_fill
            cell.border = cell_border
        
        ws.row_dimensions[row_idx].height = 22
        row_idx += 1

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def create_complete_job_zip(job):
    """Bundles entire Job digital dossier: summary, photos, QC report, ledger into a single ZIP"""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as zf:
        # 1. Job Summary Text Dossier
        summary_text = f"""========================================================================
SRI SWARNA JEWELLERS WORKSHOP - COMPLETE JOB ARCHIVE DOSSIER
========================================================================
Job ID:             {job.job_id}
Order Number:       {job.order.order_number}
Customer:           {job.order.customer.name} (Phone: {job.order.customer.mobile})
Item Description:   {job.title}
Jewellery Type:     {job.get_jewellery_type_display()}
Quantity:           {job.quantity}
Gold Purity:        {job.gold_purity}
Expected Weight:    {job.expected_weight} grams
Final Gross Weight: {job.final_gross_weight or 'N/A'} grams
Net Gold Weight:    {job.net_gold_weight or 'N/A'} grams
Stone Weight:       {job.stone_weight or 'N/A'} grams
Overall Status:     {job.get_overall_status_display()}
Current Stage:      {job.get_current_stage_display()}
Assigned Worker:    {job.current_worker.name if job.current_worker else 'None'}
Physical Location:  {job.get_current_location_display()}
Target Completion:  {job.target_date}
Date Created:       {job.created_at}

--- STAGE EXECUTION HISTORY ---
"""
        for stage in job.stage_histories.all():
            w_name = stage.worker.name if stage.worker else "Unassigned"
            summary_text += f"\n- {stage.get_stage_name_display()}: Status: {stage.get_status_display()} | Karigar: {w_name} | Duration: {stage.actual_duration_minutes}m"
            if stage.is_rework:
                summary_text += f" [REWORK: {stage.rework_reason}]"

        summary_text += f"\n\n--- GOLD ACCOUNTABILITY LEDGER ---"
        summary_text += f"\nTotal Gold Issued:   {job.gold_issued_total}g"
        summary_text += f"\nTotal Gold Returned: {job.gold_returned_total}g"
        summary_text += f"\nDifference:          {job.gold_balance_difference}g\n"
        for txn in job.gold_transactions.all():
            summary_text += f"\n  [{txn.transaction_id}] {txn.created_at.strftime('%d/%m/%Y %H:%M')}: {txn.get_transaction_type_display()} - {txn.weight_grams}g ({txn.notes})"

        # QC status
        if hasattr(job, 'quality_check'):
            qc = job.quality_check
            summary_text += f"\n\n--- QUALITY CONTROL CHECKLIST ---\nResult: {qc.get_result_display()} | Passed items: {qc.passed_items_count}/10"
            if qc.failure_reason:
                summary_text += f"\nFailure/Rework Notes: {qc.failure_reason}"

        zf.writestr(f"{job.job_id}/Job_Summary_{job.job_id}.txt", summary_text)

        # 2. Photos packaging
        for photo in job.photos.all():
            if photo.image and os.path.exists(photo.image.path):
                category_dir = "reference" if "REFERENCE" in photo.category else ("final" if "FINAL" in photo.category else "progress")
                filename = os.path.basename(photo.image.path)
                zf.write(photo.image.path, f"{job.job_id}/{category_dir}/{filename}")

        # 3. QR Code
        if job.qr_code_image and os.path.exists(job.qr_code_image.path):
            zf.write(job.qr_code_image.path, f"{job.job_id}/qr_code_{job.job_id}.png")

    buffer.seek(0)
    return buffer


def perform_system_backup(notes="Manual Backup initiated from Backup Center"):
    """Creates a full timestamped snapshot of SQLite DB, media files and creates BackupRecord"""
    backup_dir = settings.BACKUPS_DIR
    os.makedirs(backup_dir, exist_ok=True)

    timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_filename = f"GSCMS_Backup_{timestamp_str}.zip"
    backup_path = os.path.join(backup_dir, backup_filename)

    with zipfile.ZipFile(backup_path, 'w', zipfile.ZIP_DEFLATED) as zf:
        # 1. SQLite Database
        db_path = settings.DATABASES['default']['NAME']
        if os.path.exists(db_path):
            zf.write(db_path, "db.sqlite3")

        # 2. Media Directory (all job photos, QR codes, avatars)
        media_root = settings.MEDIA_ROOT
        if os.path.exists(media_root):
            for root, dirs, files in os.walk(media_root):
                for file in files:
                    full_p = os.path.join(root, file)
                    rel_p = os.path.relpath(full_p, media_root)
                    zf.write(full_p, os.path.join("media", rel_p))

    file_size = os.path.getsize(backup_path)
    record = BackupRecord.objects.create(
        backup_file=backup_filename,
        backup_type="FULL_ARCHIVE",
        file_size_bytes=file_size,
        status="SUCCESS",
        notes=notes
    )
    return record
